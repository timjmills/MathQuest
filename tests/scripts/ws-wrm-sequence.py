#!/usr/bin/env python3
"""ws-wrm-sequence.py -- build the White Rose small-step page sequence (GENERATED data).

PRIMARY source (owner 2026-10-17): the LIVE school curriculum site https://timjmills.github.io/awsajacademymath/ —
its CURRICULUM.grades (decoded from the page's <script id="__data"> block) gives each grade's domains in the site's order,
their core units and lessons in the site's order, the CCSS codes and the step flags (native / above / committee = CCSS
BUILD / non-ccss). Extracted, data only, to data/curriculum/source/awsaj-preview-sequence.json
(`--preview <saved index.html> --save-preview` refreshes it).
Secondary: the school's domain sequence workbook
  data/curriculum/source/Awsaj-Domain-Sequence-K-5-2026-27.xlsx
(one sheet per grade: Week | Qtr | Week of | Days | Unit | Domain | Lesson | CCSS | Power | Type | From | S | P | I |
 support-blocks list on the first lesson of each week) — used ONLY for what the site lacks: weeks, strands, power,
copied-from, the week's support blocks, unit weeks / power / tests, and the enrichment lessons the site does not list.
A workbook core lesson the site does not list is dropped and reported. PK comes from the site alone.

Every lesson is mapped to our WRM step id (data/curriculum/wrm-steps.json; R = PK, Y1 = K ... Y6 = Grade 5). Nothing is
guessed silently: every lesson that cannot be mapped, every WRM step the sequence leaves out and every workbook / preview
disagreement goes into the report.

Writes:
  data/curriculum/wrm-sequence.json        the full GENERATED sequence
  js/modules/wrm-sequence-db.js            the compact browser copy (never hand-edit)
  design/audit/runs/wrm-page/MAPPING_REPORT.md
Usage: python3 tests/scripts/ws-wrm-sequence.py [--preview <file>] [--save-preview] [--check]
  --check  exit 1 when the generated files differ from what is on disk (no write).
"""
import base64, gzip, json, os, re, sys

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
XLSX = os.path.join(ROOT, 'data', 'curriculum', 'source', 'Awsaj-Domain-Sequence-K-5-2026-27.xlsx')
STEPS = os.path.join(ROOT, 'data', 'curriculum', 'wrm-steps.json')
OUT_JSON = os.path.join(ROOT, 'data', 'curriculum', 'wrm-sequence.json')
OUT_JS = os.path.join(ROOT, 'js', 'modules', 'wrm-sequence-db.js')
OUT_REPORT = os.path.join(ROOT, 'design', 'audit', 'runs', 'wrm-page', 'MAPPING_REPORT.md')
# The preview's sequence, extracted (data only) so --check works without the site file.
PREVIEW_DEFAULT = os.path.join(ROOT, 'data', 'curriculum', 'source', 'awsaj-preview-sequence.json')
# Lessons the school teaches with US money where White Rose teaches pounds: the nearest current money step
# (owner ruling 2026-10-10: map them, tagged "ADAPT: US money").
MONEY_ADAPT = {
    ('1', 'find the total'): 'Y2.B4.S7',
    ('1', 'find the difference'): 'Y2.B4.S7',
    ('1', 'count money notes and coins'): 'Y2.B4.S2',
    ('1', 'select money'): 'Y2.B4.S4',
    ('3', 'pounds tens ones and pence'): 'Y4.B10.S1',
}

GRADES = [  # id, sheet, WRM year, label
    ('PK', None, 'R', 'Pre-K'),
    ('K', 'Kindergarten', 'Y1', 'Kindergarten'),
    ('1', 'Grade 1', 'Y2', 'Grade 1'),
    ('2', 'Grade 2', 'Y3', 'Grade 2'),
    ('3', 'Grade 3', 'Y4', 'Grade 3'),
    ('4', 'Grade 4', 'Y5', 'Grade 4'),
    ('5', 'Grade 5', 'Y6', 'Grade 5'),
]
YEAR_OF_GRADE = {g: y for g, _, y, _ in GRADES}
PREVIEW_GRADE = {'PK': 'PK', 'K': 'K', '1': 'G1', '2': 'G2', '3': 'G3', '4': 'G4', '5': 'G5'}
DOMAIN_CODE = {
    'Counting & Cardinality': 'CC', 'Operations & Algebraic Thinking': 'OA',
    'Number & Operations in Base Ten': 'NBT', 'Number & Operations: Fractions': 'NF',
    'Measurement & Data': 'MD', 'Geometry': 'G', 'Enrichment lessons': 'E',
}
TYPE_CODE = {'White Rose': 'wr', 'COPIED IN': 'copied', 'CCSS BUILD (to be made)': 'build', 'Enrichment (after MAP)': 'enrich'}
PRIOR_PREFIX = {'Rec/PK4': 'R', 'Y1/KG': 'Y1', 'Y2/Gr.1': 'Y2', 'Y3/Gr.2': 'Y3', 'Y4/Gr.3': 'Y4', 'Y5/Gr.4': 'Y5', 'Y6/Gr.5': 'Y6'}
FROM_YEAR = {'PK': 'R', 'K': 'Y1', 'G1': 'Y2', 'G2': 'Y3', 'G3': 'Y4', 'G4': 'Y5', 'G5': 'Y6'}
YEARS = ['R', 'Y1', 'Y2', 'Y3', 'Y4', 'Y5', 'Y6']


def norm(t):
    t = str(t or '').lower().replace('–', '-').replace('—', '-').replace('’', "'")
    t = re.sub(r'[^a-z0-9]+', ' ', t)
    t = re.sub(r'\b(\d) d\b', r'\1d', t)          # 3-D / 3D
    t = re.sub(r'\b([ap]) m\b', r'\1m', t)          # a.m. / am
    return re.sub(r'\s+', ' ', t).strip()


def loose(t):
    t = re.sub(r'\([^)]*\)', ' ', str(t or ''))
    t = norm(t).replace('recognize', 'recognise').replace('color', 'colour')
    return t


def load_steps():
    d = json.load(open(STEPS))
    steps, blocks = {}, {}
    exact, rough = {}, {}  # (year, norm) / (year, loose) -> [ids]; kept apart so an exact title always wins
    for y in d['years']:
        for b in y['blocks']:
            blocks[b['id']] = b
            for s in b['steps']:
                steps[s['id']] = dict(s, year=y['id'], block=b['id'], blockName=b['name'])
                exact.setdefault((y['id'], norm(s['title'])), []).append(s['id'])
                rough.setdefault((y['id'], loose(s['title'])), []).append(s['id'])
    return steps, blocks, (exact, rough)


def load_preview(path):
    if not path:
        return None
    raw = open(path, encoding='utf-8').read()
    if path.endswith('.json'):
        return json.loads(raw)
    m = re.search(r'<script[^>]*id="__data"[^>]*>(.*?)</script>', raw, re.S)
    return json.loads(gzip.decompress(base64.b64decode(m.group(1).strip())).decode('utf-8'))


def step_order(sid):
    y, b, s = sid.split('.')
    return (YEARS.index(y), int(b[1:]), int(s[1:]))


class Mapper:
    def __init__(self, steps, blocks, by_title, preview):
        self.steps, self.blocks, self.by_title, self.preview = steps, blocks, by_title, preview
        self.prev_units = {}  # (grade, norm title) -> [unit names]
        if preview:
            for g, pg in PREVIEW_GRADE.items():
                grade = preview['CURRICULUM']['grades'].get(pg)
                if not grade:
                    continue
                for dom in grade['domains']:
                    for u in dom.get('coreUnits', []):
                        for s in u['steps']:
                            self.prev_units.setdefault((g, norm(s['title'])), []).append(u['name'])

    def candidates(self, year, title):
        exact, rough = self.by_title
        return exact.get((year, norm(title))) or rough.get((year, loose(title))) or []

    def pick(self, cands, grade, title, ccss, near):
        """Choose among same-title steps of one year: preview unit name, then CCSS overlap, then nearest."""
        cands = sorted(set(cands), key=step_order)
        if len(cands) == 1:
            return cands[0], ''
        units = set(self.prev_units.get((grade, norm(title)), []))
        by_unit = [c for c in cands if self.steps[c]['blockName'] in units]
        if len(by_unit) == 1:
            return by_unit[0], 'preview unit'
        pool = by_unit or cands
        cc = set(ccss)
        by_cc = [c for c in pool if cc & set(self.steps[c].get('ccss') or [])]
        if len(by_cc) == 1:
            return by_cc[0], 'CCSS'
        pool = by_cc or pool
        if near:
            pool = sorted(pool, key=lambda c: abs(step_order(c)[1] - step_order(near)[1]) * 100 + abs(step_order(c)[2] - step_order(near)[2]))
        return pool[0], 'nearest of %d' % len(cands)

    def lesson(self, grade, title, ccss, typ, frm, near):
        """Map one workbook lesson. Returns (step id or None, how)."""
        home = YEAR_OF_GRADE[grade]
        if typ == 'build':
            return None, 'CCSS BUILD (no White Rose lesson)'
        adapt = MONEY_ADAPT.get((grade, loose(title))) if 'US' in str(title) else None
        if adapt:
            return adapt, 'ADAPT: US money'
        order = []
        if typ == 'copied' and frm in FROM_YEAR:
            order.append(FROM_YEAR[frm])
        order.append(home)
        # enrichment and stray titles: other years nearest first
        for d in (-1, 1, -2, 2, -3, 3, -4, 4, -5, 5, -6, 6):
            i = YEARS.index(home) + d
            if 0 <= i < len(YEARS):
                order.append(YEARS[i])
        for y in order:
            c = self.candidates(y, title)
            if c:
                sid, how = self.pick(c, grade, title, ccss, near if near and near.startswith(y + '.') else None)
                where = '' if y == home or (typ == 'copied' and y == order[0]) else ' (CROSS-YEAR from %s)' % y
                return sid, ('title' + (', ' + how if how else '') + where)
        return None, 'no WRM step with this title'

    def prior(self, text):
        """'Y1/KG Count in 10s' -> (year, title, step id or None)."""
        text = text.strip()
        for p, y in PRIOR_PREFIX.items():
            if text.startswith(p + ' '):
                title = text[len(p) + 1:].strip()
                c = self.candidates(y, title)
                sid = self.pick(c, None, title, [], None)[0] if c else None
                return y, title, sid
        return None, text, None


def split_ccss(v):
    return [c.strip() for c in str(v or '').split(',') if c.strip()]


DOMAIN_NAME = {v: k for k, v in DOMAIN_CODE.items()}
SITE_GRADE = {'PK': 'PK', 'K': 'K', '1': 'G1', '2': 'G2', '3': 'G3', '4': 'G4', '5': 'G5'}
SITE_LESSONS_JSON = os.path.join(ROOT, 'data', 'curriculum', 'site-lessons.json')


def site_lesson_ids():
    """(site grade, title slug) -> the site-lessons.json record (the school site's id for a lesson with no WRM step)."""
    out = {}
    for r in json.load(open(SITE_LESSONS_JSON, encoding='utf-8')).get('lessons', []):
        gid, _, slug = r['id'].partition('|-|')
        out[(gid, slug)] = r
    return out
KEEP_FLAGS = ('above', 'non-ccss', 'committee')


def site_order(g, pg, mp, report, label):
    """Owner 2026-10-17: the LIVE school site is the primary source of order and membership. Each site domain becomes
    one unit (D1.. in the site's domain order), its core units' steps in the site's order; the workbook gives only what
    the site lacks (weeks, strands, power, copied-from, the unit's weeks / power / test). Workbook core lessons the site
    does not list are dropped and reported; workbook enrichment lessons the site does not list stay as enrichment."""
    wb_units = g['units']
    # the workbook pass's mapping notes are re-made below for the lessons the page keeps
    for k in ('unmapped', 'ambiguous', 'crossYear'):
        report[k] = [r for r in report.get(k, []) if r[0] != label]
    site_ids = site_lesson_ids()
    by_title = {}
    for u in wb_units:
        for l in u['lessons']:
            by_title.setdefault(norm(l['title']), []).append((u, l))
    site_titles = set()
    units = []
    near = None
    # a site domain with no core lessons (Grade 1 CC: one support lesson only) gets no unit; its support lessons
    # join the nearest unit on the page (wrm-links.js placeUnit)
    doms = [d for d in pg.get('domains', []) if any(cu['steps'] for cu in d.get('coreUnits', []))]
    for d in pg.get('domains', []):
        if d not in doms:
            report.setdefault('emptyDomain', []).append((label, d['code']))
    for di, dom in enumerate(doms):
        code = dom['code']
        wbu = next((u for u in wb_units if u['domain'] == code), None)
        u = {'id': 'D%d' % (di + 1), 'domain': code, 'name': DOMAIN_NAME.get(code, dom['name']),
             'weeks': wbu['weeks'] if wbu else '', 'power': wbu['power'] if wbu else [], 'test': wbu['test'] if wbu else '',
             'lessons': []}
        for cu in dom.get('coreUnits', []):
            for st in cu['steps']:
                key = norm(st['title'])
                site_titles.add(key)
                hits = by_title.get(key, [])
                hit = next((l for wu, l in hits if wu['domain'] == code), None) or (hits[0][1] if hits else None)
                flag = st.get('flag') or ''
                cc = list(st.get('ccss') or [])
                sl = site_ids.get((SITE_GRADE[g['id']], re.sub(r'[^a-z0-9]', '', st['title'].lower())))
                if flag == 'committee':
                    sid, typ, how = None, 'build', 'CCSS BUILD'
                    if not sl:
                        report['unmapped'].append((label, u['id'], '', st['title'], flag, 'CCSS BUILD with no site-lessons.json id'))
                elif hit and hit['type'] != 'build' and hit['step']:
                    sid, typ, how = hit['step'], ('copied' if hit['type'] == 'copied' else 'wr'), hit.get('how', 'workbook')
                else:
                    sid, how = mp.lesson(g['id'], st['title'], cc, 'wr', hit.get('from') if hit else None, near)
                    typ = 'wr'
                if sid and ('CROSS-YEAR' in how or 'ADAPT' in how):
                    report.setdefault('crossYear', []).append((label, u['id'], st['title'], sid, how))
                if sid and 'nearest' in how:
                    report['ambiguous'].append((label, u['id'], st['title'], sid, how))
                if not sid and flag != 'committee':
                    if sl:
                        report.setdefault('siteId', []).append((label, u['id'], st['title'], sl['id']))
                    else:
                        report['unmapped'].append((label, u['id'], ','.join(hit['weeks']) if hit else '', st['title'], flag or 'wr', how))
                if sid:
                    near = sid
                # weeks: the workbook's; a lesson the workbook does not list takes the site's own pacing weeks
                les = {'n': len(u['lessons']) + 1, 'title': st['title'], 'step': sid, 'ccss': cc,
                       'power': bool(hit['power']) if hit else bool(sl and sl.get('power')), 'type': typ,
                       'weeks': list(hit['weeks']) if hit else list((sl or {}).get('weeks') or []),
                       'strands': hit['strands'] if hit else ''.join(k for k, t in (('S', 'standard'), ('P', 'priority'), ('I', 'intervention')) if t in ((sl or {}).get('tracks') or [])),
                       'siteUnit': cu['name']}
                if flag in KEEP_FLAGS:
                    les['flag'] = flag
                if hit and hit.get('from'):
                    les['from'] = hit['from']
                if hit and hit.get('note'):
                    les['note'] = hit['note']
                elif 'ADAPT' in how:
                    les['note'] = 'ADAPT: US money'
                if not hit:
                    report.setdefault('siteOnly', []).append((label, u['id'], st['title'], flag, ', '.join(les['weeks'])))
                u['lessons'].append(les)
        units.append(u)
    # workbook lessons the site does not list
    for wu in wb_units:
        for l in wu['lessons']:
            if norm(l['title']) in site_titles:
                continue
            if wu['id'] != 'E':
                report['disagree'].append((label, 'in workbook, not on the site (dropped)', '%s: %s' % (wu['id'], l['title'])))
    enrich = next((u for u in wb_units if u['id'] == 'E'), None)
    if enrich:
        kept = [l for l in enrich['lessons'] if norm(l['title']) not in site_titles]
        for i, l in enumerate(kept):
            l['n'] = i + 1
        if kept:
            units.append(dict(enrich, lessons=kept))
    g['units'] = units
    g['source'] = 'site (order, membership) + workbook (weeks, strands, power)'


def build(preview_path):
    steps, blocks, by_title = load_steps()
    preview = load_preview(preview_path)
    mp = Mapper(steps, blocks, by_title, preview)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    summary = {}
    for r in wb['Summary'].iter_rows(min_row=2, values_only=True):
        if r[0] and r[1]:
            summary[(r[0], r[1])] = r
    report = {'unmapped': [], 'ambiguous': [], 'priorUnmapped': [], 'disagree': [], 'missing': {}}
    grades = []
    used = set()
    for gid, sheet, year, label in GRADES:
        g = {'id': gid, 'label': label, 'year': year, 'source': 'workbook' if sheet else 'preview', 'units': [], 'weeks': {}}
        if sheet:
            ws = wb[sheet]
            units = {}
            near = None
            cur_week = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                week, qtr, weekof, days, unit, domain, title, ccss, power, typ, frm, s, p, i = row[:14]
                support = row[14] if len(row) > 14 else None
                if not week:
                    continue
                if week != cur_week:
                    cur_week = week
                    g['weeks'].setdefault(week, {'qtr': qtr, 'weekOf': weekof, 'days': days, 'prior': []})
                if support:
                    pr = []
                    for part in str(support).split(';'):
                        if not part.strip():
                            continue
                        y, t, sid = mp.prior(part)
                        if not y:   # the exam row's 15th column names the domain test, not a lesson
                            g['weeks'][week]['test'] = part.strip()
                            continue
                        pr.append({'year': y, 'title': t, 'step': sid})
                        if not sid:
                            report['priorUnmapped'].append((label, week, part.strip()))
                    g['weeks'][week]['prior'] = pr
                tc = TYPE_CODE.get(typ)
                if not tc or not unit:
                    if typ == 'EXAM':
                        g['weeks'][week]['exam'] = domain
                    if typ == 'MAP':
                        g['weeks'][week]['map'] = True
                    continue
                if unit not in units:
                    sm = summary.get((sheet, unit))
                    u = {'id': unit, 'domain': DOMAIN_CODE.get(domain, domain), 'name': domain,
                         'weeks': sm[3] if sm else '', 'power': split_ccss(sm[9]) if sm else [],
                         'test': (sm[10] if sm else '') or '', 'lessons': []}
                    units[unit] = u
                    g['units'].append(u)
                u = units[unit]
                strands = ''.join(k for k, v in zip('SPI', (s, p, i)) if v)
                key = norm(title)
                hit = next((l for l in u['lessons'] if norm(l['title']) == key), None)
                if hit:
                    if week not in hit['weeks']:
                        hit['weeks'].append(week)
                    hit['strands'] = ''.join(k for k in 'SPI' if k in hit['strands'] or k in strands)
                    continue
                cc = split_ccss(ccss)
                sid, how = mp.lesson(gid, title, cc, tc, frm, near)
                if sid:
                    near = sid
                    used.add(sid)
                    if 'nearest' in how:
                        report['ambiguous'].append((label, unit, title, sid, how))
                    if 'CROSS-YEAR' in how or 'ADAPT' in how:
                        report.setdefault('crossYear', []).append((label, unit, title, sid, how))
                elif tc != 'build':
                    report['unmapped'].append((label, unit, week, title, tc, how))
                les = {'n': len(u['lessons']) + 1, 'title': title, 'step': sid, 'ccss': cc, 'power': power == 'POWER',
                       'type': tc, 'weeks': [week], 'strands': strands}
                if frm:
                    les['from'] = frm
                if 'ADAPT' in how:
                    les['note'] = 'ADAPT: US money'
                les['how'] = how
                u['lessons'].append(les)
            # the live site orders the grade (owner 2026-10-17); the old cross-check below is kept for the workbook
            if preview and preview['CURRICULUM']['grades'].get(PREVIEW_GRADE[gid]):
                site_order(g, preview['CURRICULUM']['grades'][PREVIEW_GRADE[gid]], mp, report, label)
                report.setdefault('siteCount', {})[gid] = sum(len(cu['steps']) for d in preview['CURRICULUM']['grades'][PREVIEW_GRADE[gid]]['domains'] for cu in d.get('coreUnits', []))
            elif preview:
                pg = preview['CURRICULUM']['grades'].get(PREVIEW_GRADE[gid])
                ours = {norm(l['title']) for u in g['units'] if u['id'] != 'E' for l in u['lessons']}
                theirs = {}
                for dom in (pg or {}).get('domains', []):
                    for un in dom.get('coreUnits', []):
                        for st in un['steps']:
                            theirs[norm(st['title'])] = (dom['code'], un['name'], st['title'])
                every = {norm(l['title']) for u in g['units'] for l in u['lessons']}
                for k, v in theirs.items():
                    if k not in every:
                        report['disagree'].append((label, 'in preview, not in workbook', '%s / %s: %s' % v))
                for k in ours - set(theirs):
                    t = next(l['title'] for u in g['units'] if u['id'] != 'E' for l in u['lessons'] if norm(l['title']) == k)
                    report['disagree'].append((label, 'in workbook, not in preview', t))
        elif preview and preview['CURRICULUM']['grades'].get('PK'):
            pg = preview['CURRICULUM']['grades']['PK']
            report.setdefault('siteCount', {})['PK'] = sum(len(cu['steps']) for d in pg['domains'] for cu in d.get('coreUnits', []))
            for di, dom in enumerate(pg['domains']):
                for un in dom.get('coreUnits', []):
                    u = {'id': '%s%d' % (dom['code'], un['num']), 'domain': dom['code'], 'name': '%s: %s' % (dom['name'], un['name']),
                         'weeks': '', 'power': [], 'test': '', 'lessons': []}
                    # the Reception block of this unit: the one holding most of the unit's titles
                    score = {}
                    for st in un['steps']:
                        for x in set(mp.candidates('R', st['title'])):
                            score[steps[x]['block']] = score.get(steps[x]['block'], 0) + 1
                    for st in un['steps']:
                        c = sorted(set(mp.candidates('R', st['title'])), key=lambda x: (-score.get(steps[x]['block'], 0), step_order(x)))
                        sid = c[0] if c else None
                        if len(c) > 1 and score.get(steps[c[0]]['block'], 0) == score.get(steps[c[1]]['block'], 0):
                            report['ambiguous'].append((label, u['id'], st['title'], sid, 'Reception tie of %d' % len(c)))
                        if sid:
                            used.add(sid)
                        else:
                            report['unmapped'].append((label, u['id'], '', st['title'], 'wr', 'no Reception step with this title'))
                        u['lessons'].append({'n': len(u['lessons']) + 1, 'title': st['title'], 'step': sid, 'ccss': st.get('ccss') or [],
                                             'power': False, 'type': 'wr', 'weeks': [], 'strands': ''})
                    g['units'].append(u)
        else:
            g['note'] = 'No workbook sheet for Pre-K and no preview given: Pre-K omitted.'
        grades.append(g)
    # WRM steps the sequence leaves out (prior-learning lists count as "used" only for the report's second column)
    prior_used = {p['step'] for g in grades for w in g['weeks'].values() for p in w['prior'] if p['step']}
    used = {l['step'] for g in grades for u in g['units'] for l in u['lessons'] if l['step']}
    for g in grades:
        for u in g['units']:
            for l in u['lessons']:
                l.pop('how', None)
    for sid in sorted(steps, key=step_order):
        if sid not in used:
            report['missing'].setdefault(steps[sid]['year'], []).append((sid, steps[sid]['title'], sid in prior_used))
    return {'meta': {
        'title': 'Awsaj White Rose small-step sequence, K to Grade 5 by CCSS domain, 2026-27',
        'generatedBy': 'python3 tests/scripts/ws-wrm-sequence.py' + (' --preview <preview>' if preview else ''),
        'source': 'the live school site https://timjmills.github.io/awsajacademymath/ (PRIMARY since 2026-10-17: domain order, '
                  'units and lessons in its order, CCSS codes, step flags; extracted to data/curriculum/source/awsaj-preview-sequence.json); '
                  'data/curriculum/source/Awsaj-Domain-Sequence-K-5-2026-27.xlsx (only what the site lacks: weeks, strands, power, '
                  'copied-from, prior learning, unit weeks / power / tests, and the enrichment lessons the site does not list)',
        'fields': {'step': 'our WRM step id (null: CCSS BUILD lesson or unmapped)', 'type': 'wr | copied | build | enrich',
                   'strands': 'S = Standard, P = Priority, I = Intervention', 'weeks': 'weeks the lesson is taught',
                   'prior': "the week's support-block lessons (prior learning), mapped to WRM steps",
                   'flag': 'the site\'s step flag when it is above | non-ccss | committee (committee = CCSS BUILD)',
                   'siteUnit': 'the site\'s core unit (White Rose block) the lesson sits in'},
    }, 'grades': grades}, report


def write_report(seq, rep):
    lessons = [l for g in seq['grades'] for u in g['units'] for l in u['lessons']]
    mapped = sum(1 for l in lessons if l['step'])
    builds = sum(1 for l in lessons if l['type'] == 'build')
    prior = [p for g in seq['grades'] for w in g['weeks'].values() for p in w['prior']]
    L = ['# WRM page: sequence mapping report (GENERATED by tests/scripts/ws-wrm-sequence.py)', '',
         '| | count |', '|---|---|',
         '| Lessons in the sequence (unique per unit) | %d |' % len(lessons),
         '| Mapped to a WRM step | %d |' % mapped,
         '| CCSS BUILD lessons (no White Rose lesson exists; expected) | %d |' % builds,
         '| Unmapped (not CCSS BUILD) | %d |' % len(rep['unmapped']),
         '| Same-title steps settled by nearest block (check) | %d |' % len(rep['ambiguous']),
         '| Mapped to another year\'s step, or adapted (check) | %d |' % len(rep.get('crossYear', [])),
         '| Prior-learning entries | %d (unmapped %d) |' % (len(prior), len(rep['priorUnmapped'])),
         '| WRM steps the sequence leaves out | %d |' % sum(len(v) for v in rep['missing'].values()),
         '| Workbook / site disagreements (core units; enrichment not compared) | %d |' % len(rep['disagree']),
         '| Site lessons with no workbook row (weeks from the site\'s pacing) | %d |' % len(rep.get('siteOnly', [])),
         '| No WRM step, mapped to a school-site id (not CCSS BUILD) | %d |' % len(rep.get('siteId', [])), '']
    L += ['## Per grade: live site core lessons vs the page', '', '| Grade | site core lessons | page lessons (domain units) | CCSS BUILD | above | non-CCSS | enrichment kept |', '|---|---|---|---|---|---|---|']
    for g in seq['grades']:
        core = [l for u in g['units'] if u['id'] != 'E' for l in u['lessons']]
        en = [l for u in g['units'] if u['id'] == 'E' for l in u['lessons']]
        L.append('| %s | %s | %d | %d | %d | %d | %d |' % (g['label'], rep.get('siteCount', {}).get(g['id'], '-'), len(core),
                 sum(1 for l in core if l.get('flag') == 'committee'), sum(1 for l in core if l.get('flag') == 'above'),
                 sum(1 for l in core if l.get('flag') == 'non-ccss'), len(en)))
    L += ['']
    L += ['## Unmapped lessons', '']
    L += ['- %s %s %s: "%s" (%s) — %s' % r for r in rep['unmapped']] or ['None.']
    L += ['', '## Same-title steps settled by nearest block', '']
    L += ['- %s %s: "%s" -> %s (%s)' % r for r in rep['ambiguous']] or ['None.']
    L += ['', '## Mapped to another year\'s step, or adapted (US money)', '']
    L += ['- %s %s: "%s" -> %s (%s)' % r for r in rep.get('crossYear', [])] or ['None.']
    L += ['', '## Prior-learning entries that match no WRM step', '']
    L += ['- %s %s: %s' % r for r in rep['priorUnmapped']] or ['None.']
    L += ['', '## WRM steps the sequence leaves out', '', '"prior" = still listed as prior learning in some week.', '']
    for y in YEARS:
        v = rep['missing'].get(y, [])
        L.append('### %s (%d)' % (y, len(v)))
        L.append('')
        L += ['- %s %s%s' % (sid, t, ' (prior)' if p else '') for sid, t, p in v] or ['None.']
        L.append('')
    L += ['## Domain order (the site\'s, used by the page)', '']
    L += ['- %s: %s' % (g['label'], ' > '.join(u['domain'] for u in g['units'] if u['id'] != 'E' and not u['id'][:2].isalpha() or u['id'].startswith('D')) if g['id'] != 'PK'
          else ' > '.join(dict.fromkeys(u['domain'] for u in g['units']))) for g in seq['grades']]
    L += ['', 'The workbook ordered Grade 1 NBT > OA > G > MD, Grade 2 NBT > OA > MD > G, Grade 3 OA > NF > NBT > MD > G, Grade 4 '
          'OA > NBT > NF > G > MD and Grade 5 NBT > OA > NF > MD > G.', '']
    L += ['- %s: the site\'s %s domain has no core lessons (support lessons only), so the page has no unit for it; its support '
          'lessons join the nearest unit.' % r for r in rep.get('emptyDomain', [])] + ([''] if rep.get('emptyDomain') else [])
    L += ['## Workbook / site disagreements', '', 'The site wins: a workbook lesson the site does not list is dropped from the page. '
          'Every dropped lesson is a workbook "COPIED IN" lesson (an earlier grade\'s step taught again); the site no longer lists them.', '']
    L += ['- %s, %s: %s' % r for r in rep['disagree']] or ['None (or no preview given).']
    L += ['', '## Site lessons with no workbook row (listed in the site\'s place; weeks from the site\'s pacing)', '']
    L += ['- %s %s: "%s"%s%s' % (a, b, c, ' (%s)' % d if d else '', ', site weeks %s' % w if w else '') for a, b, c, d, w in rep.get('siteOnly', [])] or ['None.']
    L += ['', '## Lessons with no WRM step, mapped to their school-site id (site-lessons.json)', '']
    L += ['- %s %s: "%s" -> %s' % r for r in rep.get('siteId', [])] or ['None.']
    return '\n'.join(L) + '\n'


def compact_js(seq):
    steps = load_steps()[0]
    used = sorted({l['step'] for g in seq['grades'] for u in g['units'] for l in u['lessons'] if l['step']}, key=step_order)
    used += sorted({p['step'] for g in seq['grades'] for w in g['weeks'].values() for p in w['prior'] if p['step'] and p['step'] not in used}, key=step_order)
    drive = {sid: [(steps[sid].get('drive') or {}).get(k, '') for k in ('lesson', 'guide', 'video')] for sid in used}
    grades = []
    for g in seq['grades']:
        weeks = {k: [p['step'] for p in w['prior'] if p['step']] for k, w in g['weeks'].items()}
        info = {k: {kk: w[kk] for kk in ('weekOf', 'exam', 'map') if w.get(kk)} for k, w in g['weeks'].items()}
        grades.append({'id': g['id'], 'label': g['label'], 'year': g['year'], 'units': [
            {'id': u['id'], 'domain': u['domain'], 'name': u['name'], 'weeks': u['weeks'], 'power': u['power'],
             'lessons': [[l['title'], l['step'], l['ccss'], 1 if l['power'] else 0, l['type'], l['weeks'], l['strands'], l.get('from', '')]
                         + ([l.get('note', ''), l['flag']] if l.get('flag') else [l['note']] if l.get('note') else [])
                         for l in u['lessons']]} for u in g['units']], 'prior': weeks, 'weekInfo': info})
    body = ',\n'.join('    ' + json.dumps(g, ensure_ascii=False, separators=(',', ':')) for g in grades)
    return ('// wrm-sequence-db.js — GENERATED by `python3 tests/scripts/ws-wrm-sequence.py` from the school\'s\n'
            '// live curriculum site (order) and domain-sequence workbook (weeks) in data/curriculum/source/. Never hand-edit.\n'
            '// Lesson tuple: [title, wrmStepId|null, ccss[], power(0/1), type (wr|copied|build|enrich), weeks[], strands, fromGrade, note?, siteFlag?].\n'
            '// Order and membership: the live school site (domains, units, lessons in its order); weeks etc.: the workbook.\n'
            '// prior: week -> WRM step ids of that week\'s prior-learning (support block) lessons, in the workbook\'s order.\n\n'
            'export const WRM_SEQUENCE = [\n' + body + ',\n];\n\n'
            '// step id -> [lesson file id, teaching guide file id, video url] (Google Drive, the school\'s White Rose folder).\n'
            'export const WRM_STEP_FILES = ' + json.dumps(drive, ensure_ascii=False, separators=(',', ':')) + ';\n')


def main():
    args = sys.argv[1:]
    preview = args[args.index('--preview') + 1] if '--preview' in args else PREVIEW_DEFAULT
    if '--save-preview' in args:
        pv = load_preview(preview)
        keep = {'note': 'Extracted from the live school curriculum site (https://timjmills.github.io/awsajacademymath/, the '
                        '<script id="__data"> block): grade -> domain (site order) -> core units -> steps (title, CCSS, note, flag) '
                        'and each domain\'s enrich list. Data only; the PRIMARY order source of tests/scripts/ws-wrm-sequence.py.',
                'CURRICULUM': {'grades': pv['CURRICULUM']['grades']}}
        open(PREVIEW_DEFAULT, 'w', encoding='utf-8').write(json.dumps(keep, ensure_ascii=False, indent=1) + '\n')
        preview = PREVIEW_DEFAULT
    seq, rep = build(preview)
    outs = {OUT_JSON: json.dumps(seq, ensure_ascii=False, indent=1) + '\n', OUT_JS: compact_js(seq), OUT_REPORT: write_report(seq, rep)}
    if '--check' in args:
        bad = [p for p, t in outs.items() if p != OUT_REPORT and (not os.path.exists(p) or open(p, encoding='utf-8').read() != t)]
        print('ws-wrm-sequence: %s' % ('OK' if not bad else 'STALE ' + ', '.join(os.path.relpath(b, ROOT) for b in bad)))
        sys.exit(1 if bad else 0)
    os.makedirs(os.path.dirname(OUT_REPORT), exist_ok=True)
    for p, t in outs.items():
        open(p, 'w', encoding='utf-8').write(t)
    print(write_report(seq, rep).split('## Unmapped')[0])


if __name__ == '__main__':
    main()
